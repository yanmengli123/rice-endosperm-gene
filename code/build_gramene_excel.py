#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Enrich the verified GrameneOryza v8 Nipponbare export (5,707 genes).

The companion downloader (GrameneOryza_5707_宸叉牳瀹炴帴鍙鎵归噺涓嬭浇.py) stores 8
columns.  This script re-queries the same Solr endpoint with a wider field list
(PubMed xrefs, coordinates, GO/TO/pathway/domain counts, external IDs), fetches
PubMed metadata for every linked PMID, and writes an analysis-ready workbook.

Same verified query:
    q=*:* AND ((capabilities:pubs))
    fq=taxon_id:(39947013)
    rows=200, start=0..5706, sort=id asc

Outputs (in the workspace's GrameneOryza directory):
    GrameneOryza_Curated_Nipponbare_5707_full.json
    GrameneOryza_Curated_Nipponbare_5707_enriched.xlsx
    GrameneOryza_Curated_Nipponbare_5707_enriched_genes.csv
    GrameneOryza_Curated_Nipponbare_5707_publications.csv
    GrameneOryza_Curated_Nipponbare_5707_pmids.csv
    pubmed_metadata.json

Usage:
    python build_gramene_excel.py
    python build_gramene_excel.py --skip-pubmed     # no NCBI calls
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import time
from pathlib import Path
from typing import Any

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DATA_ROOT = Path(__file__).resolve().parent.parent
BASE_DIR = DATA_ROOT / "GrameneOryza"
SEARCH_URL = "https://data.gramene.org/oryza_v8/search"
QUERY = "*:* AND ((capabilities:pubs))"
FILTER = "taxon_id:(39947013)"
TAXON_ID = 39947013
EXPECTED_TOTAL = 5707
PAGE_SIZE = 200

FULL_JSON = BASE_DIR / "GrameneOryza_Curated_Nipponbare_5707_full.json"
PUBMED_JSON = BASE_DIR / "pubmed_metadata.json"
XLSX_PATH = BASE_DIR / "GrameneOryza_Curated_Nipponbare_5707_enriched.xlsx"
GENES_CSV = BASE_DIR / "GrameneOryza_Curated_Nipponbare_5707_enriched_genes.csv"
PUBS_CSV = BASE_DIR / "GrameneOryza_Curated_Nipponbare_5707_publications.csv"
PMIDS_CSV = BASE_DIR / "GrameneOryza_Curated_Nipponbare_5707_pmids.csv"

FIELDS = [
    "id",
    "taxon_id",
    "name",
    "description",
    "synonyms",
    "biotype",
    "model_rep_id",
    "model_rep_name",
    "model_rep_description",
    "model_rep_identity",
    "model_rep_taxon_id",
    "region",
    "start",
    "end",
    "strand",
    "capabilities",
    "PUBMED__xrefs",
    "GO__ancestors",
    "TO__ancestors",
    "pathways__ancestors",
    "domains__ancestors",
    "Uniprot*",
    "EntrezGene__xrefs",
    "RefSeq_dna__xrefs",
    "RefSeq_peptide__xrefs",
    "EMBL__xrefs",
    "STRING__xrefs",
    "PanOryza__xrefs",
]

HEADERS = {
    "Accept": "application/json",
    "User-Agent": "Rice-endosperm-research-GrameneOryza-export/1.0",
}

ENDOSPERM_PATTERN = re.compile(
    r"\b("
    r"endosperm\w*|aleurone|amylose|amylopectin|starch\w*|waxy|wx|"
    r"floury|chalkiness|chalky|grain filling|grain quality|"
    r"glutelin\w*|prolamin\w*|globulin\w*|storage protein\w*|"
    r"amylase\w*|sucrose synthase|adp-glucose|granule-bound"
    r")\b",
    re.IGNORECASE,
)

CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").replace("\u00ad", "")
    text = CONTROL_CHARS.sub("", text)
    return re.sub(r"\s+", " ", text).strip()


def join_list(value: Any, sep: str = "; ") -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return sep.join(clean(item) for item in value if clean(item))
    return clean(value)


def fetch_page(session: requests.Session, start: int, rows: int) -> tuple[int, list[dict[str, Any]]]:
    params = {
        "q": QUERY,
        "fq": FILTER,
        "start": start,
        "rows": rows,
        "sort": "id asc",
        "fl": ",".join(FIELDS),
        "wt": "json",
    }
    last_error: Exception | None = None
    for attempt in range(1, 5):
        try:
            response = session.get(SEARCH_URL, params=params, timeout=(15, 120))
            response.raise_for_status()
            payload = response.json()["response"]
            return int(payload["numFound"]), list(payload["docs"])
        except Exception as exc:  # noqa: BLE001 - transient API/network errors
            last_error = exc
            print(f"[warn] start={start} attempt {attempt}/4: {exc}", flush=True)
            time.sleep(2 * attempt)
    raise RuntimeError(f"page start={start} failed: {last_error}")


def fetch_all() -> list[dict[str, Any]]:
    session = requests.Session()
    session.headers.update(HEADERS)
    print(f"[info] GET {SEARCH_URL}")
    print(f"[info] q={QUERY} | fq={FILTER}")

    num_found, first_docs = fetch_page(session, 0, PAGE_SIZE)
    print(f"[ok] numFound={num_found:,} (expected {EXPECTED_TOTAL:,})")
    if num_found != EXPECTED_TOTAL:
        raise SystemExit(
            f"[error] Gramene now returns {num_found:,}, not {EXPECTED_TOTAL:,}; "
            "refusing to export a mismatched snapshot."
        )

    docs = list(first_docs)
    while len(docs) < num_found:
        start = len(docs)
        total, page = fetch_page(session, start, min(PAGE_SIZE, num_found - start))
        if total != num_found or not page:
            raise SystemExit(f"[error] unstable pagination at start={start}")
        docs.extend(page)
        print(f"[ok] {len(docs):,}/{num_found:,}")
        time.sleep(0.2)
    session.close()

    if len(docs) != EXPECTED_TOTAL:
        raise SystemExit(f"[error] got {len(docs):,} records")

    ids: list[str] = []
    for index, doc in enumerate(docs, start=1):
        gene_id = doc.get("id")
        if not gene_id:
            raise SystemExit(f"[error] record {index} has no id")
        if int(doc.get("taxon_id", -1)) != TAXON_ID:
            raise SystemExit(
                f"[error] record {index} taxon_id mismatch: {doc.get('taxon_id')}"
            )
        if "pubs" not in (doc.get("capabilities") or []):
            raise SystemExit(f"[error] record {index} lacks 'pubs' capability")
        ids.append(str(gene_id))
    if len(set(ids)) != EXPECTED_TOTAL:
        raise SystemExit(f"[error] only {len(set(ids)):,} unique ids")

    FULL_JSON.write_text(json.dumps(docs, ensure_ascii=False), encoding="utf-8")
    print(f"[ok] saved {FULL_JSON.name} ({FULL_JSON.stat().st_size / 1e6:.1f} MB)")
    return docs


EUTILS = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi"


def load_pubmed_cache() -> dict[str, dict[str, Any]]:
    if PUBMED_JSON.exists():
        try:
            return json.loads(PUBMED_JSON.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return {}
    return {}


def fetch_pubmed(pmids: list[str]) -> dict[str, dict[str, Any]]:
    cache = load_pubmed_cache()
    missing = [pmid for pmid in pmids if pmid not in cache]
    print(f"[info] PubMed IDs: {len(pmids):,} unique, {len(missing):,} need metadata")
    if not missing:
        return cache

    session = requests.Session()
    session.headers.update({"User-Agent": HEADERS["User-Agent"]})
    batch_size = 150
    for offset in range(0, len(missing), batch_size):
        batch = missing[offset : offset + batch_size]
        params = {
            "db": "pubmed",
            "id": ",".join(batch),
            "retmode": "json",
            "tool": "rice_endosperm_export",
        }
        last_error: Exception | None = None
        for attempt in range(1, 4):
            try:
                response = session.get(EUTILS, params=params, timeout=60)
                response.raise_for_status()
                raw_dir = BASE_DIR / "pubmed_raw_batches"
                raw_dir.mkdir(exist_ok=True)
                (raw_dir / f"{batch[0]}_{batch[-1]}.json").write_bytes(response.content)
                result = response.json().get("result", {})
                for pmid in batch:
                    item = result.get(pmid)
                    if not isinstance(item, dict) or item.get("error"):
                        cache[pmid] = {}
                        continue
                    authors = [clean(a.get("name")) for a in item.get("authors", []) if a.get("name")]
                    cache[pmid] = {
                        "title": clean(item.get("title")),
                        "journal": clean(item.get("source")),
                        "pubdate": clean(item.get("pubdate")),
                        "volume": clean(item.get("volume")),
                        "issue": clean(item.get("issue")),
                        "pages": clean(item.get("pages")),
                        "doi": next(
                            (clean(entry.get("value")) for entry in item.get("articleids", [])
                             if entry.get("idtype") == "doi"), ""
                        ),
                        "authors": authors,
                        "pubtype": join_list(item.get("pubtype")),
                    }
                last_error = None
                break
            except Exception as exc:  # noqa: BLE001
                last_error = exc
                time.sleep(2 * attempt)
        if last_error:
            print(f"[warn] PubMed batch at {offset} failed: {last_error}")
        print(f"[ok] PubMed {min(offset + batch_size, len(missing)):,}/{len(missing):,}")
        time.sleep(0.4)

    PUBMED_JSON.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
    print(f"[ok] saved {PUBMED_JSON.name}")
    return cache



def build_tables(
    docs: list[dict[str, Any]], pubmed: dict[str, dict[str, Any]]
) -> dict[str, pd.DataFrame]:
    gene_rows: list[dict[str, Any]] = []
    publication_rows: list[dict[str, Any]] = []

    for index, doc in enumerate(docs, start=1):
        gene_id = clean(doc.get("id"))
        symbol = clean(doc.get("name") or doc.get("model_rep_name"))
        pmids = [clean(p) for p in (doc.get("PUBMED__xrefs") or []) if clean(p)]
        pmids = list(dict.fromkeys(pmids))
        numeric_pmids = [pmid for pmid in pmids if pmid.isdigit()]
        other_refs = [pmid for pmid in pmids if not pmid.isdigit()]

        gene_rows.append(
            {
                "No.": index,
                "Gene_ID": gene_id,
                "Gene_Symbol": symbol,
                "Description": clean(doc.get("description")),
                "Synonyms": join_list(doc.get("synonyms")),
                "No. of PMIDs": len(numeric_pmids),
                "PMIDs": "; ".join(numeric_pmids),
                "Biotype": clean(doc.get("biotype")),
                "Chromosome": clean(doc.get("region")),
                "Start": doc.get("start", ""),
                "End": doc.get("end", ""),
                "Strand": doc.get("strand", ""),
                "Model_Rep_ID": clean(doc.get("model_rep_id")),
                "Model_Rep_Name": clean(doc.get("model_rep_name")),
                "Model_Rep_Description": clean(doc.get("model_rep_description")),
                "Model_Rep_Identity": doc.get("model_rep_identity", ""),
                "Model_Rep_Taxon_ID": doc.get("model_rep_taxon_id", ""),
                "Capabilities": join_list(doc.get("capabilities")),
                "GO terms": len(doc.get("GO__ancestors") or []),
                "TO terms": len(doc.get("TO__ancestors") or []),
                "Pathway terms": len(doc.get("pathways__ancestors") or []),
                "Domain terms": len(doc.get("domains__ancestors") or []),
                "UniProt": join_list(doc.get("Uniprot/SWISSPROT__xrefs")),
                "UniProt gene name": join_list(doc.get("Uniprot_gn__xrefs")),
                "EntrezGene": join_list(doc.get("EntrezGene__xrefs")),
                "RefSeq mRNA": join_list(doc.get("RefSeq_dna__xrefs")),
                "RefSeq protein": join_list(doc.get("RefSeq_peptide__xrefs")),
                "EMBL": join_list(doc.get("EMBL__xrefs")),
                "STRING": join_list(doc.get("STRING__xrefs")),
                "PanOryza": join_list(doc.get("PanOryza__xrefs")),
                "Gene URL": f"https://ensembl.gramene.org/Oryza_sativa/Gene/Summary?g={gene_id}",
                "Source": "GrameneOryza v8; capabilities:pubs",
                "No. of other refs": len(other_refs),
                "Other publication refs": " | ".join(other_refs),
            }
        )

        for pmid in numeric_pmids:
            meta = pubmed.get(pmid, {})
            pubdate = clean(meta.get("pubdate"))
            year_match = re.search(r"(19|20)\d{2}", pubdate)
            publication_rows.append(
                {
                    "Gene_ID": gene_id,
                    "Gene_Symbol": symbol,
                    "PMID": pmid,
                    "Title": clean(meta.get("title")),
                    "Journal": clean(meta.get("journal")),
                    "Year": year_match.group(0) if year_match else "",
                    "First_Author": (meta.get("authors") or [""])[0],
                    "DOI": clean(meta.get("doi")),
                    "PubMed_URL": f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/",
                }
            )

    genes_df = pd.DataFrame(gene_rows)

    pmids_df = (
        pd.DataFrame(publication_rows)
        .groupby("PMID", as_index=False)
        .agg(
            Title=("Title", "first"),
            Journal=("Journal", "first"),
            Year=("Year", "first"),
            No_of_genes=("Gene_ID", "nunique"),
        )
        .sort_values("PMID", key=lambda s: pd.to_numeric(s, errors="coerce"))
    ) if publication_rows else pd.DataFrame(
        columns=["PMID", "Title", "Journal", "Year", "No_of_genes"]
    )

    return {
        "genes": genes_df,
        "publications": pd.DataFrame(publication_rows),
        "pmids": pmids_df,
    }





def find_endosperm_hits(
    genes: pd.DataFrame, publications: pd.DataFrame
) -> pd.DataFrame:
    titles: dict[str, list[str]] = {}
    if not publications.empty:
        for row in publications.to_dict("records"):
            titles.setdefault(row["Gene_ID"], []).append(str(row.get("Title") or ""))

    hits: list[dict[str, Any]] = []
    for row in genes.to_dict("records"):
        text = " ".join(
            [
                str(row.get("Gene_Symbol") or ""),
                str(row.get("Description") or ""),
                str(row.get("Synonyms") or ""),
            ]
            + titles.get(row["Gene_ID"], [])
        )
        tokens = {match.lower() for match in ENDOSPERM_PATTERN.findall(text)}
        if not tokens:
            continue
        hits.append(
            {
                "Gene_ID": row["Gene_ID"],
                "Gene_Symbol": row["Gene_Symbol"],
                "Description": row["Description"],
                "Matched keywords": "; ".join(sorted(tokens)),
                "No. of PMIDs": row["No. of PMIDs"],
                "PMIDs": row["PMIDs"],
                "Gene URL": row["Gene URL"],
            }
        )
    return pd.DataFrame(hits)


def build_readme(meta: dict[str, Any]) -> pd.DataFrame:
    rows = [
        ("GrameneOryza v8 - Nipponbare genes with publication annotation", ""),
        ("Source API", SEARCH_URL),
        ("Query", QUERY),
        ("Filter", FILTER),
        ("Taxon", "Oryza sativa japonica Nipponbare (taxon_id 39947013)"),
        ("Retrieved (local time)", meta["retrieved"]),
        ("Genes", meta["genes"]),
        ("Gene-PMID links", meta["links"]),
        ("Unique PMIDs", meta["pmids"]),
        ("Genes with >=1 PMID", meta["with_pmids"]),
        ("PMIDs without PubMed metadata", meta["missing_meta"]),
        ("Other publication refs (DOI/citation)", meta["other_refs"]),
        ("Genes with description", meta["with_desc"]),
        ("Endosperm keyword hits", meta["hits"]),
        ("", ""),
        (
            "Relationship to the original downloader",
            "GrameneOryza_Curated_Nipponbare_5707.tsv / _raw.json are unchanged "
            "outputs of the verified 200-row pagination script. This workbook adds "
            "PUBMED__xrefs, coordinates, model-representative protein, external IDs "
            "and GO/TO/pathway/domain counts from the same Solr query.",
        ),
        (
            "Publication annotations",
            "PUBMED__xrefs lists every PubMed ID attached to the gene by Gramene's "
            "curation. Titles/journal/year come from NCBI E-utilities esummary "
            "(pubmed_metadata.json is the normalized cache; pubmed_raw_batches/ preserves "
            "the exact response bytes). DOI is taken from articleids with idtype=doi.",
        ),
        (
            "GO / TO / pathway / domain counts",
            "Solr returns ancestor ontology IDs only (GO__ancestors, TO__ancestors, "
            "pathways__ancestors, domains__ancestors); the workbook reports their "
            "counts. Map IDs to term names via OBO ontologies if needed.",
        ),
        (
            "Gene URL",
            "https://ensembl.gramene.org/Oryza_sativa/Gene/Summary?g=<Gene_ID>",
        ),
        ("Script", "build_gramene_excel.py (python build_gramene_excel.py)"),
        ("Raw enriched payload", FULL_JSON.name),
    ]
    return pd.DataFrame(rows, columns=["Item", "Value"])



GENE_WIDTHS = [
    5, 15, 18, 60, 35, 8, 30, 14, 10, 10, 10, 7, 14, 16, 45, 10, 12, 45, 8, 8,
    8, 8, 14, 16, 12, 16, 16, 28, 25, 20, 50, 40, 8, 55,
]

HEADER_FILL = PatternFill("solid", fgColor="1F6B3B")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
LINK_FONT = Font(color="0563C1", underline="single", size=10)


def style_sheet(
    worksheet: Any,
    widths: list[float],
    wrap_columns: tuple[int, ...] = (),
    link_columns: tuple[int, ...] = (),
) -> None:
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 30
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    )
    worksheet.freeze_panes = "A2"
    for column in wrap_columns:
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row=row, column=column).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
    for column in link_columns:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=column)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.font = LINK_FONT


def write_outputs(
    tables: dict[str, pd.DataFrame], hits: pd.DataFrame, readme: pd.DataFrame
) -> None:
    genes = tables["genes"]
    publications = tables["publications"]
    pmids = tables["pmids"]

    genes.to_csv(GENES_CSV, index=False, encoding="utf-8-sig")
    publications.to_csv(PUBS_CSV, index=False, encoding="utf-8-sig")
    pmids.to_csv(PMIDS_CSV, index=False, encoding="utf-8-sig")

    with pd.ExcelWriter(XLSX_PATH, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        genes.to_excel(writer, sheet_name="Genes", index=False)
        publications.to_excel(writer, sheet_name="Publications", index=False)
        pmids.to_excel(writer, sheet_name="Unique PMIDs", index=False)
        hits.to_excel(writer, sheet_name="Endosperm keyword hits", index=False)

        book = writer.book
        style_sheet(book["README"], [36, 110], wrap_columns=(2,))
        style_sheet(
            book["Genes"],
            GENE_WIDTHS,
            wrap_columns=(4, 5, 7, 18, 34),
            link_columns=(31,),
        )
        style_sheet(
            book["Publications"],
            [15, 18, 12, 70, 28, 7, 14, 26, 42],
            wrap_columns=(4,),
            link_columns=(9,),
        )
        style_sheet(book["Unique PMIDs"], [12, 80, 30, 7, 10], wrap_columns=(2,))
        style_sheet(
            book["Endosperm keyword hits"],
            [15, 18, 70, 26, 8, 30, 50],
            wrap_columns=(3,),
            link_columns=(7,),
        )


def validate_outputs(expected_rows: int) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(XLSX_PATH, read_only=True)
    try:
        for name in workbook.sheetnames:
            sheet = workbook[name]
            print(f"[check] sheet '{name}': {sheet.max_row:,} rows x {sheet.max_column} cols")
        data_rows = workbook["Genes"].max_row - 1
        if data_rows != expected_rows:
            raise SystemExit(
                f"[error] 'Genes' has {data_rows:,} rows, expected {expected_rows:,}"
            )
    finally:
        workbook.close()
    print(f"[ok] workbook verified: {XLSX_PATH}")



def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Build the enriched GrameneOryza Nipponbare publication-gene workbook."
    )
    parser.add_argument(
        "--refresh-full",
        action="store_true",
        help="re-query the Solr API even if the enriched JSON already exists",
    )
    parser.add_argument(
        "--skip-pubmed",
        action="store_true",
        help="do not call NCBI E-utilities; keep PMIDs but leave titles empty",
    )
    args = parser.parse_args(argv)

    started = dt.datetime.now()
    if FULL_JSON.exists() and not args.refresh_full:
        docs = json.loads(FULL_JSON.read_text(encoding="utf-8"))
        print(f"[info] using cached {FULL_JSON.name}: {len(docs):,} records")
        if len(docs) != EXPECTED_TOTAL:
            raise SystemExit(f"[error] cached JSON has {len(docs):,} records")
    else:
        docs = fetch_all()

    ids = {str(doc["id"]) for doc in docs}
    if len(ids) != EXPECTED_TOTAL:
        raise SystemExit(f"[error] unique gene ids: {len(ids):,}")

    pmids = sorted(
        {
            str(pmid)
            for doc in docs
            for pmid in (doc.get("PUBMED__xrefs") or [])
            if str(pmid).strip().isdigit()
        },
        key=int,
    )
    other_refs_total = sum(
        1
        for doc in docs
        for pmid in (doc.get("PUBMED__xrefs") or [])
        if str(pmid).strip() and not str(pmid).strip().isdigit()
    )
    print(f"[ok] unique PubMed IDs across genes: {len(pmids):,}")
    print(f"[ok] other publication ref entries (raw): {other_refs_total:,}")

    pubmed: dict[str, dict[str, Any]] = {}
    if not args.skip_pubmed:
        pubmed = fetch_pubmed(pmids)
    else:
        print("[info] --skip-pubmed: publication titles will be empty")

    tables = build_tables(docs, pubmed)
    hits = find_endosperm_hits(tables["genes"], tables["publications"])

    genes = tables["genes"]
    publications = tables["publications"]
    missing_meta = sum(
        1 for pmid in pmids if not (pubmed.get(pmid) or {}).get("title")
    )
    meta = {
        "retrieved": dt.datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S UTC%z"),
        "genes": len(genes),
        "links": len(publications),
        "pmids": len(pmids),
        "with_pmids": int((genes["No. of PMIDs"] > 0).sum()),
        "missing_meta": missing_meta,
        "with_desc": int(genes["Description"].astype(str).str.strip().ne("").sum()),
        "hits": len(hits),
        "other_refs": int(genes["No. of other refs"].sum()),
    }
    readme = build_readme(meta)
    write_outputs(tables, hits, readme)
    validate_outputs(len(genes))

    print("\n=== Summary ===")
    print(f"Genes                 : {meta['genes']:,}")
    print(f"Gene-PMID links       : {meta['links']:,}")
    print(f"Unique PMIDs          : {meta['pmids']:,}")
    print(f"Genes with >=1 PMID   : {meta['with_pmids']:,}")
    print(f"PMIDs w/o metadata    : {meta['missing_meta']:,}")
    print(f"Endosperm keyword hits: {meta['hits']:,}")
    print(f"Elapsed               : {(dt.datetime.now() - started).total_seconds():.1f} s")
    print(f"\nWorkbook : {XLSX_PATH}")
    print(f"Genes CSV: {GENES_CSV}")
    print(f"Pubs CSV : {PUBS_CSV}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
